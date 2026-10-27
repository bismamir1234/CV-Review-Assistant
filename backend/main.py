import asyncio
import json
import os
import re
from typing import Any, Optional

from dotenv import load_dotenv

load_dotenv()

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel, Field
from sqlalchemy import DateTime, Integer, Text, func
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker
from sqlalchemy import create_engine
from groq import Groq

# ---------- Configuration ----------
DATABASE_URL = os.getenv("DATABASE_URL", "sqlite:///./cv_assistant.db")
MAX_UPLOAD_BYTES = 5 * 1024 * 1024  # 5 MB
GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GROQ_MODEL = os.getenv("GROQ_MODEL", "llama-3.3-70b-versatile")

engine = create_engine(
    DATABASE_URL,
    connect_args={"check_same_thread": False} if DATABASE_URL.startswith("sqlite") else {},
)
SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)


class Base(DeclarativeBase):
    pass


class Review(Base):
    __tablename__ = "reviews"

    id: Mapped[int] = mapped_column(Integer, primary_key=True, index=True)
    filename: Mapped[str] = mapped_column(Text, nullable=False)
    extracted_text: Mapped[str] = mapped_column(Text, nullable=False)
    job_description: Mapped[str] = mapped_column(Text, nullable=False, default="")
    result_json: Mapped[str] = mapped_column(Text, nullable=False)
    created_at: Mapped[Any] = mapped_column(DateTime(timezone=True), server_default=func.now())


Base.metadata.create_all(bind=engine)

app = FastAPI(title="CV Assistant API", version="1.2.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["http://localhost:5173", "http://127.0.0.1:5173"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)


# ---------- Pydantic Schemas ----------

class CVSuggestion(BaseModel):
    id: str = Field(..., description="Unique suggestion identifier")
    category: str = Field(..., description="Category (Grammar, Action Verbs, Clarity, Impact, etc.)")
    original_statement: str = Field(..., description="Original statement from CV")
    issue_identified: str = Field(..., description="Explanation of identified issue")
    suggested_rewrite: str = Field(..., description="Proposed truthful revision preserving original facts")
    explanation: str = Field(..., description="Why this improvement helps")
    facts_to_verify: str = Field(..., description="Facts, dates, metrics, or tools the user must verify")
    status: str = Field(default="pending", description="Status: pending, accepted, rejected, edited")


class CVReviewRequest(BaseModel):
    cv_text: str
    target_role: Optional[str] = ""


class CVReviewResponse(BaseModel):
    summary: str
    suggestions: list[CVSuggestion]
    mode: str  # "ai" or "demo"


class SkillMatch(BaseModel):
    skill: str
    evidence: str


class JobMatchRequest(BaseModel):
    cv_text: str
    job_description: str


class JobMatchResponse(BaseModel):
    matched_skills: list[SkillMatch]
    evidence_summary: str
    possible_gaps: list[str]
    relevant_keywords: list[str]
    supported_keywords: list[str]
    verification_keywords: list[str]
    uncertain_matches: list[str]
    recommendations: list[str]
    matching_score: int
    score_explanation: str
    mode: str  # "ai" or "demo"


class Analysis(BaseModel):
    summary: str
    suggestions: list[CVSuggestion]
    revised_text: str
    job_match: Optional[JobMatchResponse] = None
    mode: str  # "ai" or "demo"


class ReviewResponse(BaseModel):
    id: int
    filename: str
    extracted_text: str
    job_description: str
    analysis: Analysis


# ---------- Helper Functions & PII Redaction ----------

def sanitize_cv_for_ai(text: str) -> str:
    """Redact unnecessary personal identifying information (PII) before sending to AI."""
    # Redact email addresses
    sanitized = re.sub(r'\b[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Z|a-z]{2,}\b', '[REDACTED EMAIL]', text)
    # Redact phone numbers
    sanitized = re.sub(r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}', '[REDACTED PHONE]', sanitized)
    # Redact postal codes
    sanitized = re.sub(r'\b\d{5}(?:-\d{4})?\b', '[REDACTED POSTAL]', sanitized)
    return sanitized


def extract_docx_text(data: bytes) -> str:
    """Extract normal paragraphs and table-cell text from a DOCX file."""
    from io import BytesIO

    try:
        document = Document(BytesIO(data))
    except Exception as exc:
        raise HTTPException(
            status_code=400,
            detail="The DOCX file could not be opened. It may be damaged or password-protected.",
        ) from exc

    chunks: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                chunks.append(line)

    extracted = "\n".join(chunks).strip()
    if not extracted:
        raise HTTPException(
            status_code=400,
            detail="No readable text was found in this DOCX. Please upload a text-based CV.",
        )
    return extracted


# ---------- Demo Engine (Local Heuristic Fallback) ----------

def demo_review_suggestions(cv_text: str, target_role: str = "") -> CVReviewResponse:
    """Local heuristic suggestions when GROQ_API_KEY is not set or unavailable."""
    lines = [line.strip() for line in cv_text.splitlines() if line.strip()]
    suggestions: list[CVSuggestion] = []

    weak_phrases = [
        ("responsible for", "Spearheaded / Managed delivery of", "Action Verbs & Ownership",
         "Passive phrasing describes duties rather than ownership and measurable impact.",
         "Verify whether you led this initiative independently or collaborated with a team."),
        ("worked on", "Architected / Engineered / Deployed", "Technical Leadership",
         "'Worked on' obscures your exact engineering contribution.",
         "Verify the exact components or modules you personally designed or engineered."),
        ("helped with", "Partnered cross-functionally to execute", "Professional Wording",
         "Minimizes your active collaboration and proactive contribution.",
         "Confirm your specific deliverable and partner teams involved."),
        ("assisted in", "Collaborated on production release of", "Active Contribution",
         "Sounds like passive observation rather than accountable execution.",
         "Verify the technical milestones you helped achieve."),
        ("duties included", "Delivered core operational milestones including", "Impact & Clarity",
         "Focuses on job description responsibilities instead of achievements.",
         "Verify the specific outcomes or metrics generated by these duties."),
    ]

    sugg_counter = 1
    for line in lines:
        lower = line.lower()
        for phrase, replacement, cat, issue, verify in weak_phrases:
            if phrase in lower and len(suggestions) < 4:
                idx = lower.find(phrase)
                rewritten = line[:idx] + replacement + line[idx + len(phrase):]
                suggestions.append(CVSuggestion(
                    id=f"sugg-{sugg_counter}",
                    category=cat,
                    original_statement=line[:250],
                    issue_identified=issue,
                    suggested_rewrite=rewritten[:300],
                    explanation=f"Replaces '{phrase}' with strong, proactive wording to highlight candidate initiative.",
                    facts_to_verify=verify,
                    status="pending",
                ))
                sugg_counter += 1
                break

    # Check for quantification opportunities
    metric_pattern = re.compile(r'(\d+[\d,.]*[%$kKmMbB]?|\$\d+)')
    has_metrics = any(metric_pattern.search(line) for line in lines)
    if not has_metrics and len(suggestions) < 5:
        for line in lines:
            if len(line) > 35 and any(term in line.lower() for term in ["system", "api", "project", "feature", "data", "app"]):
                suggestions.append(CVSuggestion(
                    id=f"sugg-{sugg_counter}",
                    category="Quantified Impact",
                    original_statement=line[:250],
                    issue_identified="Lacks quantifiable scale, volume, or performance metrics.",
                    suggested_rewrite=line + " (e.g., resulting in [X]% latency reduction or supporting [N] daily active users).",
                    explanation="Quantifying achievements using Google's XYZ formula makes engineering achievements verifiable and impactful.",
                    facts_to_verify="Only add real metrics from production, testing, or analytics. Do not fabricate figures.",
                    status="pending",
                ))
                sugg_counter += 1
                break

    # Clarity & Readability check
    for line in lines:
        if len(line.split()) > 35 and len(suggestions) < 6:
            suggestions.append(CVSuggestion(
                id=f"sugg-{sugg_counter}",
                category="Sentence Clarity & Brevity",
                original_statement=line[:250],
                issue_identified="Sentence is over 35 words long and may fatigue recruiters or ATS readers.",
                suggested_rewrite="Break this statement into two concise bullet points: first lead with the core achievement, followed by the technologies used.",
                explanation="Shorter sentences under 25 words improve scanning speed and comprehension for hiring managers.",
                facts_to_verify="Ensure all original factual claims remain intact when splitting the sentence.",
                status="pending",
            ))
            sugg_counter += 1
            break

    if not suggestions:
        first_line = lines[0] if lines else "Resume Header"
        suggestions.append(CVSuggestion(
            id=f"sugg-{sugg_counter}",
            category="Professional Formatting",
            original_statement=first_line[:200],
            issue_identified="Header alignment and ATS metadata check.",
            suggested_rewrite=first_line + " | LinkedIn | GitHub | Portfolio",
            explanation="Clean header layout guarantees automated parsers extract contact details without errors.",
            facts_to_verify="Verify your profile URLs are current and functional.",
            status="pending",
        ))

    summary = (
        f"Demo Review (Local Heuristic Engine): Processed {len(lines)} lines. "
        f"Identified {len(suggestions)} practical wording, clarity, and impact improvements. "
        f"(Configure GROQ_API_KEY in backend/.env for full deep LLM analysis)."
    )
    return CVReviewResponse(summary=summary, suggestions=suggestions, mode="demo")


def demo_match_job(cv_text: str, job_description: str) -> JobMatchResponse:
    """Local heuristic job matching when GROQ_API_KEY is not set or unavailable."""
    if not job_description.strip():
        raise HTTPException(status_code=400, detail="Job description cannot be empty for job matching.")

    cv_lower = cv_text.lower()
    jd_lower = job_description.lower()

    # Common technical terms, frameworks, tools & certifications
    tech_keywords = [
        "python", "javascript", "typescript", "react", "vue", "angular", "node", "fastapi",
        "django", "flask", "sql", "postgresql", "mysql", "mongodb", "sqlite", "redis",
        "aws", "azure", "gcp", "docker", "kubernetes", "git", "ci/cd", "linux", "rest",
        "graphql", "html", "css", "tailwind", "machine learning", "pytorch", "tensorflow",
        "agile", "scrum", "microservices", "unit testing", "pytest", "terraform", "cloud",
        "solutions architect", "security", "devops", "java", "c++", "c#"
    ]

    # Find keywords present in the JD
    jd_found_keywords = [kw for kw in tech_keywords if kw in jd_lower]
    if not jd_found_keywords:
        # Fallback keyword extraction: words >= 4 letters
        words = re.findall(r'\b[a-zA-Z]{4,}\b', job_description)
        stopwords = {"with", "that", "this", "from", "have", "will", "your", "must", "their", "about", "which", "would", "there", "experience", "skills", "ability", "years", "work"}
        jd_found_keywords = [w.lower() for w in words if w.lower() not in stopwords][:12]

    matched_skills: list[SkillMatch] = []
    possible_gaps: list[str] = []
    supported_keywords: list[str] = []
    verification_keywords: list[str] = []

    # Find evidence in CV
    cv_sentences = re.split(r'[\n.!?]+', cv_text)
    for kw in jd_found_keywords:
        # Check if keyword is present but negated (e.g. "no cloud certifications", "without AWS")
        is_negated = any(re.search(rf'\b(?:no|not|without|never|lack)\s+(?:\w+\s+)?{re.escape(kw)}\b', s.lower()) for s in cv_sentences)
        if kw in cv_lower and not is_negated:
            supported_keywords.append(kw)
            # Find matching sentence
            ev = next((s.strip() for s in cv_sentences if kw in s.lower() and len(s.strip()) > 10), f"Mentioned in CV text: '{kw}'")
            matched_skills.append(SkillMatch(skill=kw.title(), evidence=ev[:250]))
        else:
            verification_keywords.append(kw)
            possible_gaps.append(
                f"[Not found in provided CV] {kw.title()} — Note: This indicates the skill was not found in your uploaded text, not that you lack the capability."
            )

    # Uncertain matches
    uncertain = []
    for soft in ["leadership", "cross-functional collaboration", "agile", "mentorship", "system design"]:
        if soft in jd_lower and soft not in cv_lower:
            uncertain.append(f"'{soft.title()}' requirement in job description may be partially evidenced through team projects; verify if you have relevant experience.")

    # Recommendations
    recs = []
    if verification_keywords:
        top_gaps = [k.title() for k in verification_keywords[:3]]
        recs.append(f"If you possess genuine experience with {', '.join(top_gaps)}, consider explicitly detailing these technologies in your projects or skills section.")
    recs.append("Ensure every technical skill listed in your CV is backed by an actual bullet point describing how you applied it.")
    recs.append("Tailor your professional summary to highlight the core domain requirements of the target role.")

    total_reqs = len(jd_found_keywords) or 1
    match_score = min(100, max(15, round((len(supported_keywords) / total_reqs) * 100)))

    return JobMatchResponse(
        matched_skills=matched_skills[:8],
        evidence_summary=f"Found {len(matched_skills)} direct requirement matches in the provided CV text.",
        possible_gaps=possible_gaps[:8],
        relevant_keywords=jd_found_keywords[:15],
        supported_keywords=supported_keywords[:10],
        verification_keywords=verification_keywords[:10],
        uncertain_matches=uncertain[:4],
        recommendations=recs,
        matching_score=match_score,
        score_explanation="This score measures keyword and requirement coverage between the provided CV and job description. It is NOT a hiring prediction, ATS pass guarantee, or assessment of your qualifications.",
        mode="demo"
    )


# ---------- Groq AI Engine Integration ----------

def _groq_review_suggestions(cv_text: str, target_role: str = "") -> CVReviewResponse:
    client = Groq(api_key=GROQ_API_KEY)
    sanitized_cv = sanitize_cv_for_ai(cv_text[:16000])

    system_prompt = """
You are a rigorous, truthful AI CV Writing Assistant.
Review the CV text thoroughly for:
1. Grammar and spelling.
2. Sentence clarity and readability.
3. Professional wording.
4. Relevant action verbs.
5. Relevance of CV content to the target role.
6. Specific, practical improvement suggestions.

STRICT SAFETY AND FACTUALITY RULES:
- NEVER invent qualifications, certifications, degrees, skills, job titles, dates, achievements, experience, or performance metrics.
- Preserve the original factual meaning.
- Clearly distinguish existing facts from proposed wording.
- Return for each suggestion:
  * "id": e.g. "sugg-1"
  * "category": one of ["Grammar & Spelling", "Sentence Clarity", "Professional Wording", "Action Verbs", "Role Relevance", "Measurable Impact"]
  * "original_statement": exact or short excerpt from CV
  * "issue_identified": specific problem identified
  * "suggested_rewrite": proposed improved phrasing preserving facts
  * "explanation": why this rewrite improves the CV
  * "facts_to_verify": specific metrics, tools, or dates the user MUST verify before accepting

Output strictly valid JSON with this schema:
{
  "summary": "overview of CV review findings",
  "suggestions": [
    {
      "id": "sugg-1",
      "category": "...",
      "original_statement": "...",
      "issue_identified": "...",
      "suggested_rewrite": "...",
      "explanation": "...",
      "facts_to_verify": "..."
    }
  ]
}
"""
    user_prompt = f"TARGET ROLE:\n{target_role.strip() or '(Not specified)'}\n\nCV TEXT (PII-Sanitized):\n{sanitized_cv}"

    completion = client.chat.completions.create(
        model=GROQ_MODEL,
        temperature=0.2,
        max_tokens=3000,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    )

    raw = completion.choices[0].message.content or "{}"
    parsed = json.loads(raw)

    raw_suggs = parsed.get("suggestions", [])
    if not isinstance(raw_suggs, list):
        raw_suggs = []

    clean_suggs: list[CVSuggestion] = []
    for idx, item in enumerate(raw_suggs[:8], 1):
        if not isinstance(item, dict):
            continue
        clean_suggs.append(CVSuggestion(
            id=str(item.get("id") or f"sugg-{idx}"),
            category=str(item.get("category", "General Improvement"))[:80],
            original_statement=str(item.get("original_statement", item.get("original", "")))[:1000],
            issue_identified=str(item.get("issue_identified", item.get("reason", "")))[:1000],
            suggested_rewrite=str(item.get("suggested_rewrite", item.get("suggested", "")))[:1500],
            explanation=str(item.get("explanation", item.get("reason", "")))[:1000],
            facts_to_verify=str(item.get("facts_to_verify", "Verify that all stated dates, tools, and achievements are accurate."))[:1000],
            status="pending"
        ))

    return CVReviewResponse(
        summary=str(parsed.get("summary", "CV Review completed with actionable suggestions."))[:2000],
        suggestions=clean_suggs,
        mode="ai"
    )


def _groq_match_job(cv_text: str, job_description: str) -> JobMatchResponse:
    client = Groq(api_key=GROQ_API_KEY)
    sanitized_cv = sanitize_cv_for_ai(cv_text[:16000])

    system_prompt = """
You are a precise, honest Job Description and CV Matching Assistant.
Compare the uploaded CV with the target job description.

RULES:
1. Identify matched skills and requirements, quoting specific evidence from the CV.
2. Identify possible gaps where job requirements are NOT evidenced in the CV.
   CRITICAL: Explain that the skill was not found in the provided CV; DO NOT assume that the candidate lacks the ability.
3. Extract relevant keywords from the job description.
4. Categorize keywords into:
   - supported_keywords: evidenced by actual experience in CV.
   - verification_keywords: required by job description but not in CV; user must verify before use.
5. Identify uncertain or ambiguous matches where claims are vague.
6. Provide honest recommendations for improving the CV for that role.
7. NEVER invent or attribute unsupported skills, certifications, or years of experience.
8. Calculate a keyword & requirement coverage score (0-100). Do NOT present this as a hiring probability.

Return strictly valid JSON with this schema:
{
  "matched_skills": [{"skill": "...", "evidence": "..."}],
  "evidence_summary": "overview of matched requirements",
  "possible_gaps": ["Not found in provided CV: ..."],
  "relevant_keywords": ["..."],
  "supported_keywords": ["..."],
  "verification_keywords": ["..."],
  "uncertain_matches": ["..."],
  "recommendations": ["..."],
  "matching_score": 75,
  "score_explanation": "This score measures keyword and requirement coverage based on the provided CV text; it is not a hiring or interview probability prediction."
}
"""
    user_prompt = f"JOB DESCRIPTION:\n{job_description[:8000]}\n\nCV TEXT (PII-Sanitized):\n{sanitized_cv}"

    completion = client.chat.completions.create(
        model=GROQ_MODEL,
        temperature=0.2,
        max_tokens=3500,
        response_format={"type": "json_object"},
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
    )

    raw = completion.choices[0].message.content or "{}"
    parsed = json.loads(raw)

    raw_matches = parsed.get("matched_skills", [])
    clean_matches: list[SkillMatch] = []
    if isinstance(raw_matches, list):
        for item in raw_matches[:12]:
            if isinstance(item, dict):
                clean_matches.append(SkillMatch(
                    skill=str(item.get("skill", ""))[:100],
                    evidence=str(item.get("evidence", ""))[:500]
                ))

    return JobMatchResponse(
        matched_skills=clean_matches,
        evidence_summary=str(parsed.get("evidence_summary", "Review the matched skills below."))[:1000],
        possible_gaps=[str(g)[:400] for g in parsed.get("possible_gaps", []) if isinstance(g, str)][:10],
        relevant_keywords=[str(k)[:80] for k in parsed.get("relevant_keywords", []) if isinstance(k, str)][:20],
        supported_keywords=[str(k)[:80] for k in parsed.get("supported_keywords", []) if isinstance(k, str)][:15],
        verification_keywords=[str(k)[:80] for k in parsed.get("verification_keywords", []) if isinstance(k, str)][:15],
        uncertain_matches=[str(u)[:400] for u in parsed.get("uncertain_matches", []) if isinstance(u, str)][:6],
        recommendations=[str(r)[:500] for r in parsed.get("recommendations", []) if isinstance(r, str)][:6],
        matching_score=int(parsed.get("matching_score", 70)),
        score_explanation=str(parsed.get("score_explanation", "This score measures keyword and requirement coverage based on the provided CV text. It is NOT a hiring prediction.")),
        mode="ai"
    )


# ---------- API Endpoints ----------

@app.get("/api/health")
def health():
    return {"status": "ok", "ai_configured": bool(GROQ_API_KEY)}


@app.post("/api/cv/review", response_model=CVReviewResponse)
async def review_cv_endpoint(payload: CVReviewRequest):
    """Review CV text for grammar, clarity, professional wording, action verbs, and actionable suggestions."""
    cv_text = payload.cv_text.strip()
    if not cv_text:
        raise HTTPException(status_code=400, detail="CV text cannot be empty.")

    if not GROQ_API_KEY:
        return demo_review_suggestions(cv_text, payload.target_role or "")

    try:
        return await asyncio.to_thread(_groq_review_suggestions, cv_text, payload.target_role or "")
    except Exception as exc:
        raise HTTPException(
            status_code=502,
            detail=f"AI review failed: {str(exc)[:100]}. Check your server API key or network.",
        ) from exc


@app.post("/api/cv/match-job", response_model=JobMatchResponse)
async def match_job_endpoint(payload: JobMatchRequest):
    """Compare CV text with a job description to identify matched skills, gaps, keywords, and recommendations."""
    cv_text = payload.cv_text.strip()
    job_description = payload.job_description.strip()

    if not cv_text:
        raise HTTPException(status_code=400, detail="CV text cannot be empty.")
    if not job_description:
        raise HTTPException(status_code=400, detail="Job description cannot be empty for job matching.")

    if not GROQ_API_KEY:
        return demo_match_job(cv_text, job_description)

    try:
        return await asyncio.to_thread(_groq_match_job, cv_text, job_description)
    except Exception as exc:
        raise HTTPException(
            status_code=502,
            detail=f"Job matching failed: {str(exc)[:100]}. Check your server API configuration.",
        ) from exc


@app.post("/api/reviews", response_model=ReviewResponse)
async def create_review(
    cv_file: UploadFile = File(...),
    job_description: str = Form(default=""),
):
    """Multipart upload endpoint combining DOCX extraction, AI review, optional job matching, and SQLite storage."""
    filename = cv_file.filename or "uploaded_cv.docx"
    if not filename.lower().endswith(".docx"):
        raise HTTPException(status_code=415, detail="MVP supports DOCX files only.")
    data = await cv_file.read(MAX_UPLOAD_BYTES + 1)
    if len(data) > MAX_UPLOAD_BYTES:
        raise HTTPException(status_code=413, detail="File is too large. Maximum size is 5 MB.")
    if not data:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")

    cv_text = extract_docx_text(data)

    # 1. Review suggestions
    if GROQ_API_KEY:
        try:
            review_resp = await asyncio.to_thread(_groq_review_suggestions, cv_text, job_description[:500])
        except Exception:
            review_resp = demo_review_suggestions(cv_text, job_description[:500])
    else:
        review_resp = demo_review_suggestions(cv_text, job_description[:500])

    # 2. Optional Job Matching
    job_match: Optional[JobMatchResponse] = None
    if job_description.strip():
        if GROQ_API_KEY:
            try:
                job_match = await asyncio.to_thread(_groq_match_job, cv_text, job_description.strip())
            except Exception:
                job_match = demo_match_job(cv_text, job_description.strip())
        else:
            job_match = demo_match_job(cv_text, job_description.strip())

    analysis = Analysis(
        summary=review_resp.summary,
        suggestions=review_resp.suggestions,
        revised_text=cv_text,
        job_match=job_match,
        mode=review_resp.mode
    )

    # Store in local SQLite
    db = SessionLocal()
    try:
        row = Review(
            filename=filename[:255],
            extracted_text=cv_text,
            job_description=job_description[:12000],
            result_json=json.dumps(analysis.model_dump(), ensure_ascii=False),
        )
        db.add(row)
        db.commit()
        db.refresh(row)
        return ReviewResponse(
            id=row.id,
            filename=row.filename,
            extracted_text=row.extracted_text,
            job_description=row.job_description,
            analysis=analysis,
        )
    finally:
        db.close()


@app.delete("/api/reviews/{review_id}")
def delete_review(review_id: int):
    db = SessionLocal()
    try:
        row = db.get(Review, review_id)
        if row is None:
            raise HTTPException(status_code=404, detail="Review not found.")
        db.delete(row)
        db.commit()
        return {"deleted": True}
    finally:
        db.close()
